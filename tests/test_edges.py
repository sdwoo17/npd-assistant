"""Bounded, reproducible property and failure tests. No live provider calls."""
import base64
import csv
import io
import json
import os
import random
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from app.contracts import date_value, validate_answer
from app.ingest import decode_file, parse_csv, chunks, redact
from app.store import AppError, Store
from app.voc import inferred_filters, source_key
from tests.helpers import Fixture, encoded, CANARY

SEEDS = tuple(range(20))


class EdgeTests(unittest.TestCase):
    def setUp(self):
        self.f = Fixture()
        self.s, self.p = self.f.service, self.f.po

    def tearDown(self):
        self.f.close()

    def test_twenty_seeds_mixed_csv_rows_and_dedupe(self):
        texts = ['소재 리포트 비교', 'Creative experiment report', '素材の比較とレポート',
                 'Werbemittel Signifikanz', 'Créatif et expérience', 'comma, quote " and\nnewline', '🧪 café ＡＢ']
        for seed in SEEDS:
            rng = random.Random(seed)
            buf = io.StringIO(); writer = csv.writer(buf)
            writer.writerow(['external_id', 'text', 'occurred_at', 'evidence_type', 'segment'])
            valid = set(); errors = 0
            for i in range(50):
                rid = str(rng.randrange(35))
                value = rng.choice(texts) + ' contact user@example.test 010-1234-5678'
                date = '2024-02-29' if rng.random() > .2 else '2025-02-29'
                kind = rng.choice(['real', 'synthetic', 'synthetic'])
                if date == '2025-02-29': errors += 1
                else: valid.add(rid)
                writer.writerow([rid, value, date, kind, 'seg-' + str(i % 3)])
            with self.subTest(seed=seed):
                result = self.s.voc_upload(self.f.owner, encoded('loop.csv', buf.getvalue(), source_name=f'seed-{seed}'))
                self.assertEqual(result['imported'], len(valid))
                self.assertEqual(len(result['errors']), errors)
                self.assertEqual(result['duplicates'], 50 - errors - len(valid))
                again = self.s.voc_upload(self.f.owner, encoded('loop.csv', buf.getvalue(), source_name=f'seed-{seed}'))
                self.assertEqual(again['imported'], 0)
                self.assertEqual(again['duplicates'], 50 - errors)
                public = json.dumps(self.s.knowledge(self.p['project_id']))
                for secret in (CANARY, 'user@example.test', '010-1234-5678'):
                    self.assertNotIn(secret, public)

    def test_twenty_seeds_stateful_publication_history(self):
        # Each sequence uses independent DB state; exact versions must remain authoritative.
        for seed in SEEDS:
            f = Fixture(); rng = random.Random(seed)
            try:
                for step in range(15):
                    old = f.store.get('project-a', 'insight', f.insight['id'])
                    operation = rng.choice(['release', 'revoke', 'edit', 'chat', 'cross_project'])
                    with self.subTest(seed=seed, step=step, operation=operation):
                        if operation in ('release', 'revoke'):
                            f.service.insight_release(f.owner, {'insight_id': old['id'], 'published': operation == 'release', 'expected_version': old['version']})
                        elif operation == 'edit':
                            f.service.insight_save(f.owner, {**old, 'insight_id': old['id'], 'expected_version': old['version'], 'text': '소재 분석 리포트 조건을 확인하는 가설.'})
                        elif operation == 'chat':
                            conv = f.conversation()
                            f.service.post(f.po, '/api/chat', {'conversation_id': conv['id'], 'message': '소재 분석 리포트', 'request_id': f'{seed}-{step}'})
                        else:
                            with self.assertRaises(AppError):
                                f.store.get('project-b', 'insight', old['id'])
                        current = f.store.get('project-a', 'insight', old['id'])
                        visible = {r['id']: r for r in f.service.knowledge('project-a')}
                        self.assertEqual(old['id'] in visible, current['published'])
                        for conv in f.store.list('project-a', 'conversation'):
                            package = f.service.export_package(f.po, conv['id'])
                            self.assertNotIn(CANARY, json.dumps(package))
                            for m in package['conversation']['messages']:
                                if not m.get('redacted'):
                                    self.assertTrue(f.service.accessible('project-a', m))
                        history = f.store.history('project-a', 'insight', old['id'])
                        self.assertEqual([r['version'] for r in history], list(range(1, current['version'] + 1)))
            finally:
                f.close()

    def test_date_boundaries_exact_range_and_invalid_types(self):
        for value in ('2024-02-29', '2026-09-12T10:11:12+09:00', '2026-09-12T01:11:12.123Z'):
            self.assertEqual(date_value(value), value)
        for value in ('2025-02-29', '2026-13-01', '2026-09-12T24:00:00Z', '2026-09-12T10:00:00', '2026-1-1', '2026-09-12T10:00:00+25:00', True, []):
            with self.subTest(value=value), self.assertRaises(AppError):
                date_value(value)
        r = inferred_filters('2026-08-19부터 2026-09-05까지 분석')
        self.assertEqual((r['date_from'], r['date_to']), ('2026-08-19', '2026-09-05'))
        self.assertEqual(inferred_filters('2024年2月')['date_to'], '2024-02-29')
        with self.assertRaises(AppError): inferred_filters('2026-09-10부터 2026-08-01')

    def test_csv_header_and_capacity_boundaries(self):
        for content in ('', 'external_id,text,text\n1,a,b', 'text\na', 'external_id,text\n1,"unterminated'):
            with self.subTest(content=content), self.assertRaises(AppError): parse_csv(content)
        rows, errors = parse_csv('external_id,text\n1,ok,extra\n2,valid')
        self.assertEqual((len(rows), len(errors)), (1, 1))
        good = 'external_id,text\n' + ''.join(f'{i},소재 분석\n' for i in range(1000))
        self.assertEqual(len(parse_csv(good)[0]), 1000)
        with self.assertRaises(AppError): parse_csv(good + '1001,extra\n')

    def test_file_decoding_docx_pdf_and_resource_limits(self):
        from docx import Document
        from pypdf import PdfWriter
        doc = Document(); doc.add_paragraph('한글 소재 분석'); doc.add_table(rows=1, cols=1).cell(0, 0).text = '표 내용'
        buf = io.BytesIO(); doc.save(buf)
        r = decode_file({'filename': 'source.docx', 'content_base64': base64.b64encode(buf.getvalue()).decode()})
        self.assertIn('표 내용', r[1]); self.assertIn('한글', r[1])
        for count in (1, 101):
            pdf = PdfWriter()
            for _ in range(count): pdf.add_blank_page(width=72, height=72)
            buf = io.BytesIO(); pdf.write(buf)
            with self.assertRaises(AppError): decode_file({'filename': 'scan.pdf', 'content_base64': base64.b64encode(buf.getvalue()).decode()})
        for size in (0, 150001):
            with self.assertRaises(AppError): decode_file(encoded('big.txt', 'x' * size))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', compression=zipfile.ZIP_DEFLATED) as z: z.writestr('word/document.xml', 'a' * (21 * 1024 * 1024))
        with self.assertRaises(AppError): decode_file({'filename': 'bomb.docx', 'content_base64': base64.b64encode(buf.getvalue()).decode()})
        self.assertEqual(len(decode_file(encoded('limit.txt', '한' * 150000))[1]), 150000)
        for seed in SEEDS:
            rng = random.Random(seed)
            text = ''.join(rng.choice('한a \n🧪') for _ in range(rng.randrange(1, 70000)))
            self.assertEqual(''.join(chunks(text)), text)
            self.assertTrue(all(len(c) <= 12000 for c in chunks(text)))

    def test_quantitative_units_and_statistics_scope(self):
        e = {'id': 'e', 'text': '25건의 합성 의견', 'evidence_type': 'synthetic'}
        a = {'text': '25건', 'evidence_ids': ['e'], 'assumptions': []}
        self.assertEqual(validate_answer(a, [e])['text'], '25건')
        for invalid in ('25%', '25명', '25원', '2026건'):
            with self.assertRaises(AppError): validate_answer({**a, 'text': invalid}, [e], {'computed_at': '2026-09-12', 'total': 100, 'denominator': 100})
        self.assertEqual(validate_answer({**a, 'text': '50% · 50건'}, [e], {'total': 100, 'denominator': 100, 'counts': {'a': 50}})['text'], '50% · 50건')

    def test_empty_research_extract_kind_is_rejected_before_model(self):
        with self.assertRaises(AppError): self.s.research_extract(self.f.owner, {'source_id': self.f.source['id'], 'evidence_type': 'real'})
        self.assertEqual(self.f.model.calls, [])

    def test_public_rss_continues_past_filtered_first_page(self):
        def page(rid, date):
            return {'feed': {'entry': [{'id': {'label': rid}, 'title': {'label': '리포트'}, 'content': {'label': '소재 분석'}, 'updated': {'label': date}, 'im:rating': {'label': '3'}}]}}
        calls = []
        def fetch(url, token):
            self.assertIn('/sortBy=mostRecent/json', url)
            calls.append(url)
            return page('first', '2026-09-01') if 'page=1/' in url else page('second', '2026-08-01') if 'page=2/' in url else {'feed': {'entry': []}}
        with patch.object(self.s, 'fetch_review_page', side_effect=fetch):
            r = self.s.collect_reviews(self.f.owner, {'provider': 'apple_public_rss', 'app_id': '123', 'max_pages': 3, 'filters': {'date_from': '2026-08-01', 'date_to': '2026-08-31'}})
        self.assertEqual((r['imported'], r['pages'], r['has_more']), (1, 3, False))
        record = self.s.voc_records('project-a')[0]
        self.assertEqual(record['external_id'], 'second')
        self.assertEqual(record['evidence_type'], 'real')
        self.assertIn('광고주 여부 미확인', record['segment'])

    def test_managed_reviews_pagination_ssrf_loop_partial_failure_retry(self):
        base = 'https://api.appstoreconnect.apple.com/v1/apps/123/customerReviews?limit=100&sort=-createdDate'
        def data(next_url=None, rid='first'):
            return {'data': [{'id': rid, 'attributes': {'title': '소재', 'body': '리포트', 'createdDate': '2026-08-01T00:00:00Z', 'rating': 4}}], 'links': {'next': next_url}}
        for next_url in ('http://127.0.0.1/', 'https://attacker.example/', 'https://api.appstoreconnect.apple.com:8443/x', 'https://token@api.appstoreconnect.apple.com/x', base):
            called = []
            def fetch(url, token): called.append(url); return data(next_url)
            with self.subTest(url=next_url), patch.dict(os.environ, {'APP_STORE_CONNECT_TOKEN': 'test-only-token'}), patch.object(self.s, 'fetch_review_page', side_effect=fetch):
                with self.assertRaises(AppError): self.s.collect_reviews(self.f.owner, {'app_id': '123', 'max_pages': 2})
            self.assertEqual(len(called), 1)
            job = self.f.store.list('project-a', 'job')[-1]
            self.assertEqual(job['status'], 'failed')
            self.assertNotIn('test-only-token', json.dumps(job))
            with patch.dict(os.environ, {'APP_STORE_CONNECT_TOKEN': 'test-only-token'}), patch.object(self.s, 'fetch_review_page', return_value=data()):
                result = self.s.retry_job(self.f.owner, {'job_id': job['id']})
            self.assertEqual(result['imported'], 0)
            self.assertEqual(self.f.store.get('project-a', 'job', job['id'])['attempts'], 2)

    def test_model_configuration_badge_requires_success(self):
        # Fixture has no login membership; bootstrap uses its trusted server-side test scope.
        self.assertFalse(self.s.get(self.p, '/api/bootstrap')['model_connection_verified'])
        c = self.f.conversation()
        self.s.post(self.p, '/api/chat', {'conversation_id': c['id'], 'message': '소재 리포트'})
        self.assertTrue(self.s.get(self.p, '/api/bootstrap')['model_connection_verified'])

    def test_database_migration_preserves_old_rows_and_project_session(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            uid = store.create_user('legacy@example.test', 'Legacy-safe-pass!', 'po', 'legacy')
            token, _ = store.login('legacy@example.test', 'Legacy-safe-pass!')
            with store.db() as db:
                db.execute('ALTER TABLE sessions DROP COLUMN project_id')
                db.execute('DROP TABLE memberships'); db.execute('DROP TABLE projects')
                db.execute('DROP TABLE record_history'); db.execute('DROP TABLE project_state')
                record = {'id': 'legacy-record', 'project_id': 'legacy', 'kind': 'prd', 'title': '기존 기획', 'created_at': '2026-01-01'}
                db.execute('INSERT INTO records VALUES(?,?,?,?,?)', ('legacy-record', 'legacy', 'prd', json.dumps(record), record['created_at']))
            upgraded = Store(tmp)
            user = upgraded.authenticate(token)
            self.assertEqual((user['project_id'], user['role']), ('legacy', 'po'))
            self.assertEqual(upgraded.get('legacy', 'prd', 'legacy-record')['version'], 1)
            self.assertEqual(upgraded.projects(uid)[0]['id'], 'legacy')

    def test_interrupted_job_can_be_retried_after_single_server_restart(self):
        job = self.s.new_job(self.f.owner, 'research_extract', {'source_id': self.f.source['id']})
        self.assertEqual(self.f.store.recover_jobs(), 1)
        self.assertEqual(self.f.store.recover_jobs(), 0)
        self.assertEqual(self.f.store.get('project-a', 'job', job['id'])['status'], 'failed')
        self.s.retry_job(self.f.owner, {'job_id': job['id']})
        self.assertEqual(self.f.store.get('project-a', 'job', job['id'])['status'], 'completed')

    def test_synthetic_insights_are_not_labeled_actual_research_grounding(self):
        self.f.store.update('project-a', 'insight', self.f.insight['id'], {'evidence_type': 'synthetic'})
        p = self.f.persona()
        self.assertEqual(p['grounding_counts']['research'], 0)
        self.assertEqual(p['grounding_counts']['synthetic_research'], 1)
        self.assertEqual(p['grounding_status'], 'synthetic_or_partial_evidence')

    def test_demo_seeds_features_voc_baseline_and_three_evidence_linked_personas(self):
        from manage import seed
        with tempfile.TemporaryDirectory() as tmp, patch('builtins.print'):
            store = Store(tmp); seed(store)
            project = 'synthetic-advertiser-project'
            self.assertEqual(len(store.list(project, 'voc')), 30)
            self.assertEqual(len(store.list(project, 'prd')), 1)
            people = store.list(project, 'persona')
            self.assertEqual(len(people), 3)
            for person in people:
                self.assertEqual(person['grounding_status'], 'synthetic_or_partial_evidence')
                self.assertEqual(person['grounding_counts']['real_voc'], 0)
                self.assertEqual({e['kind'] for e in person['evidence_snapshots']}, {'insight', 'voc'})


class BedrockSDKTests(unittest.TestCase):
    def test_real_botocore_request_contract_with_stubbed_transport(self):
        import boto3
        from botocore.stub import Stubber, ANY
        from app.model import BedrockModel
        client = boto3.client('bedrock-runtime', region_name='us-east-1', aws_access_key_id='TEST', aws_secret_access_key='TEST')
        with patch.dict(os.environ, {'BEDROCK_MODEL_ID': 'test-model', 'BEDROCK_OUTPUT_MODE': 'json_schema'}):
            model = BedrockModel()
        model.client = client
        # Both the SDK request shape and response envelope are checked by botocore.
        response = {'output': {'message': {'role': 'assistant', 'content': [{'text': '{"keywords": ["소재"]}'}]}},
                    'stopReason': 'end_turn', 'usage': {'inputTokens': 1, 'outputTokens': 1, 'totalTokens': 2}, 'metrics': {'latencyMs': 1}}
        with Stubber(client) as stub:
            stub.add_response('converse', response, {'modelId': 'test-model', 'system': ANY, 'messages': ANY, 'inferenceConfig': ANY, 'outputConfig': ANY})
            self.assertEqual(model.generate('search', {'query': '소재'})['keywords'], ['소재'])
            stub.assert_no_pending_responses()
