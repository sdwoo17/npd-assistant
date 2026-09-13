"""Pilot edge cases; all research, credentials and model responses are synthetic."""
import base64
import io
import json
import os
import random
import tempfile
import threading
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from cryptography.fernet import Fernet
from app.assets import register_templates
from app.demo_assets import initialize_pack
from app.ingest import decode_file, retrieve
from app.operations import backup, restore, existing_store, provision_account, change_account, MAGIC, backup_cipher
from app.server import Server
from app.service import Service
from app.store import Store, AppError
from scripts.https_validation import https_origin, run_remote
from scripts.register_persona_templates import register
from scripts.evaluate_retrieval import evaluate
from tests.helpers import Fixture, RecordingModel, CANARY, encoded
from tests.test_demo_assets import fixture_pack, CREDENTIALS


class RetrievalAndDocuments(unittest.TestCase):
    def test_evaluation_reports_measured_misses_and_excludes_private_text(self):
        f = Fixture()
        try:
            result = evaluate(f.service, 'project-a', [
                {'query': '소재 리포트', 'expected_evidence_ids': [f.insight['id']]},
                {'query': 'zzzz-never-match', 'expected_evidence_ids': [f.insight['id']]}], passes=3)
            self.assertEqual(result['hit_rate_at_24'], 0.5)
            self.assertTrue(result['all_ranks_stable'])
            self.assertEqual(result['model_calls'], 0)
            self.assertNotIn(CANARY, json.dumps(result))
            self.assertNotIn('zzzz-never-match', json.dumps(result))
            with self.assertRaises(AppError): evaluate(f.service, 'project-a', [{'query': 'unknown', 'insight_keys': ['missing']}])
        finally:
            f.close()

    def test_actual_claim_beats_long_metadata_stuffed_with_query_terms(self):
        rows = [{'id': 'good', 'text': '할인율이 다른 기간의 소재 비교는 판단을 보류한다.'}]
        rows += [{'id': 'noise-' + str(i), 'text': '월간 광고비 다운로드 ' * 600,
                  'feature_terms': '할인율 소재 비교 판단 보류 기간'} for i in range(80)]
        self.assertEqual(retrieve(rows, '할인율이 바뀐 소재 비교의 판단 보류', 5)[0]['id'], 'good')

    def test_channel_volume_cannot_bury_research_and_revoked_sources_never_return(self):
        f = Fixture()
        try:
            csv = 'external_id,text,evidence_type\n' + ''.join(f'{i},소재 리포트 비교 조건 불편,synthetic\n' for i in range(120))
            f.service.voc_upload(f.owner, encoded('test.csv', csv, source_name='synthetic'))
            rows, _ = f.service.search('project-a', '소재 리포트 비교')
            self.assertIn(f.insight['id'], {e['id'] for e in rows})
            f.service.insight_release(f.owner, {'insight_id': f.insight['id'], 'published': False})
            rows, _ = f.service.search('project-a', '소재 리포트 비교')
            self.assertNotIn(f.insight['id'], {e['id'] for e in rows})
            self.assertNotIn(CANARY, json.dumps(rows))
            self.assertEqual(f.service.search('other-project', '소재 리포트 비교')[0], [])
        finally:
            f.close()

    def test_repeatable_rank_loop_does_not_depend_on_record_order(self):
        rows = [{'id': str(i), 'text': text} for i, text in enumerate([
            '소재 실험 승인', '광고비 리포트', '식품 할인율 소재 비교', '신규 구매자 타깃', '재고 부족'])]
        for query in ['소재 실험', 'creative experiment', '割引 素材', 'Werbemittel', '', 'xyz-no-match']:
            expected = [r['id'] for r in retrieve(rows, query)]
            for seed in range(20):
                shuffled = rows[:]
                random.Random(seed).shuffle(shuffled)
                self.assertEqual([r['id'] for r in retrieve(shuffled, query)], expected)

    def test_docx_keeps_table_between_its_surrounding_paragraphs(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('before table')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'inside table'
        doc.add_paragraph('after table')
        stream = io.BytesIO(); doc.save(stream)
        text = decode_file({'filename': 'order.docx', 'content_base64': base64.b64encode(stream.getvalue()).decode()})[1]
        self.assertLess(text.index('before'), text.index('inside'))
        self.assertLess(text.index('inside'), text.index('after'))

    def pptx(self, *, external=False, entity=False, empty=False):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as z:
            z.writestr('ppt/presentation.xml', '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId r:id="second"/><p:sldId r:id="first"/></p:sldIdLst></p:presentation>')
            z.writestr('ppt/_rels/presentation.xml.rels', '<Relationships><Relationship Id="first" Target="slides/slide1.xml"/><Relationship Id="second" Target="slides/slide2.xml"' + (' TargetMode="External"' if external else '') + '/></Relationships>')
            for n in (1, 2):
                z.writestr(f'ppt/slides/slide{n}.xml', ('<!DOCTYPE doc [<!ENTITY x "hidden">]>' if entity else '') + '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:p><a:r><a:t>' + ('' if empty else f'visible-{n}') + '</a:t></a:r></a:p></p:sld>')
        return {'filename': 'presentation.pptx', 'content_base64': base64.b64encode(stream.getvalue()).decode()}

    def test_pptx_order_and_malformed_inputs(self):
        text = decode_file(self.pptx())[1]
        self.assertLess(text.index('visible-2'), text.index('visible-1'))
        for options in ({'external': True}, {'entity': True}, {'empty': True}):
            with self.subTest(options=options), self.assertRaises(AppError):
                decode_file(self.pptx(**options))


class DeferredProfiles(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.pack = self.root / 'pack'; fixture_pack(self.pack)
        self.model = RecordingModel()
        self.receipt = initialize_pack(self.pack, self.root / 'data', self.model, CREDENTIALS, False)
        self.store = Store(self.root / 'data'); self.s = Service(self.store, self.model)
        self.users = {}
        for role in CREDENTIALS:
            token, _ = self.store.login(*CREDENTIALS[role]); self.users[role] = self.store.authenticate(token)
        self.owner, self.po = self.users['owner'], self.users['po']; self.p = self.po['project_id']

    def tearDown(self):
        self.tmp.cleanup()

    def release(self):
        self.s.insight_release(self.owner, {'insight_ids': list(self.receipt['insight_ids'].values()), 'published': True})

    def body(self):
        rows = self.s.template_previews(self.owner)
        return {'template_ids': [r['id'] for r in rows], 'versions': {r['id']: r['version'] for r in rows}}

    def test_private_then_release_activate_retry_and_keep_po_edits(self):
        self.assertTrue(all(r['status'] == 'blocked' for r in self.s.template_previews(self.owner)))
        with self.assertRaises(AppError):
            self.s.activate_templates(self.owner, self.body())
        self.assertEqual(self.store.list(self.p, 'persona'), [])
        self.release(); body = self.body()
        result = self.s.post(self.owner, '/api/assets/activate-personas', body)
        self.assertEqual(result['created'], 8)
        person = result['personas'][0]
        self.s.save_persona(self.po, {**person, 'persona_id': person['id'], 'expected_version': person['version'], 'goals': 'PO가 수정한 목표'})
        again = self.s.activate_templates(self.owner, body)
        self.assertEqual(again['created'], 0)
        self.assertEqual(again['personas'][0]['goals'], 'PO가 수정한 목표')
        self.assertNotIn(CANARY, json.dumps(self.s.get(self.po, '/api/bootstrap')))
        self.assertEqual(self.model.calls, [])
        conv = self.s.create_conversation(self.po, {'title': '활성화 뒤 FGI', 'mode': 'interview', 'persona_ids': [r['id'] for r in result['personas']]})
        response = self.s.chat(self.po, {'conversation_id': conv['id'], 'message': '소재 리포트 비교 조건은?'})
        self.assertEqual(sum(m.get('is_synthetic', False) for m in response['messages']), 8)

    def test_owner_scope_stale_versions_duplicates_and_cross_project(self):
        for route in ('/api/assets/persona-templates', '/api/assets/activate-personas'):
            with self.assertRaises(AppError) as exc:
                self.s.get(self.po, route) if 'templates' in route else self.s.post(self.po, route, self.body())
            self.assertEqual(exc.exception.status, 403)
        self.release(); body = self.body()
        invalid = [dict(body, versions={}), dict(body, template_ids=[body['template_ids'][0]] * 2)]
        for payload in invalid:
            with self.assertRaises(AppError): self.s.activate_templates(self.owner, payload)
        with self.assertRaises(AppError):
            self.s.activate_templates({**self.owner, 'project_id': 'other'}, body)
        self.assertEqual(self.store.list(self.p, 'persona'), [])

    def test_concurrent_activation_and_midflight_revocation_are_atomic(self):
        self.release(); body = self.body()
        def call(_):
            try: return self.s.activate_templates(self.owner, body)['created']
            except AppError as exc: return exc.status
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(call, range(8)))
        self.assertEqual(results.count(8), 1)
        self.assertEqual(len(self.store.list(self.p, 'persona')), 8)
        self.s.insight_release(self.owner, {'insight_ids': list(self.receipt['insight_ids'].values()), 'published': False})
        with self.assertRaises(AppError): self.s.activate_templates(self.owner, body)
        self.assertEqual(self.s.get(self.po, '/api/bootstrap')['personas'], [])

    def test_activation_rolls_back_all_profiles_on_alias_conflict_or_revocation(self):
        self.release(); body = self.body()
        first = self.store.get(self.p, 'persona_template', body['template_ids'][0])
        definition = json.loads(self.store.decrypt(first['encrypted_definition']))
        self.s.save_persona(self.po, definition)
        with self.assertRaises(AppError): self.s.activate_templates(self.owner, body)
        self.assertEqual(len(self.store.list(self.p, 'persona')), 1)
        self.assertTrue(all(r.get('persona_id') is None for r in self.store.list(self.p, 'persona_template')))
        original = self.s.save_persona
        revoked = False
        def racing(*args, **kwargs):
            nonlocal revoked
            if not revoked:
                revoked = True
                self.s.insight_release(self.owner, {'insight_ids': list(self.receipt['insight_ids'].values()), 'published': False})
            return original(*args, **kwargs)
        with patch.object(self.s, 'save_persona', side_effect=racing), self.assertRaises(AppError):
            self.s.activate_templates(self.owner, body)
        self.assertEqual(len(self.store.list(self.p, 'persona')), 1)

    def test_legacy_registration_is_idempotent_and_rejects_changed_pack(self):
        with self.store.db() as db: db.execute("DELETE FROM records WHERE kind='persona_template'")
        first = register(self.s, self.owner, self.pack, self.receipt)
        self.assertEqual(len(first), 8)
        self.assertEqual(register(self.s, self.owner, self.pack, self.receipt), first)
        (self.pack / 'research.md').write_text('changed private source')
        with self.assertRaises(AppError): register(self.s, self.owner, self.pack, self.receipt)
        self.assertEqual(len(self.store.list(self.p, 'persona_template')), 8)


class OperatorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.store = Store(self.root / 'data')
        self.store.create_user('owner@test.example', 'Owner-initial-private!', 'owner', 'p')
        self.store.create_user('po@test.example', 'Planner-initial-private!', 'po', 'p')

    def tearDown(self): self.tmp.cleanup()

    def test_operator_and_validation_entrypoints_import_without_shadowing_stdlib(self):
        import subprocess
        import sys
        for script in ('pilot_ops.py', 'register_persona_templates.py', 'evaluate_retrieval.py',
                       'https_validation.py', 'live_validation.py', 'regression.py', 'start_bedrock.py'):
            with self.subTest(script=script):
                result = subprocess.run([sys.executable, 'scripts/' + script, '--help'], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)

    def test_account_lifecycle_revokes_sessions_and_preserves_last_owner(self):
        token, _ = self.store.login('po@test.example', 'Planner-initial-private!')
        change_account(self.store, 'po@test.example', password='Planner-rotated-private!')
        with self.assertRaises(AppError): self.store.authenticate(token)
        with self.assertRaises(AppError): self.store.login('po@test.example', 'Planner-initial-private!')
        token, _ = self.store.login('po@test.example', 'Planner-rotated-private!')
        change_account(self.store, 'po@test.example', disabled=True)
        with self.assertRaises(AppError): self.store.authenticate(token)
        with self.assertRaises(AppError): self.store.login('po@test.example', 'Planner-rotated-private!')
        change_account(self.store, 'po@test.example', disabled=False)
        self.store.login('po@test.example', 'Planner-rotated-private!')
        with self.assertRaises(AppError): change_account(self.store, 'owner@test.example', disabled=True)
        provision_account(self.store, 'second@test.example', 'Second-private-password!', 'owner', 'p')
        change_account(self.store, 'owner@test.example', disabled=True)
        with self.assertRaises(AppError): change_account(self.store, 'second@test.example', disabled=True)

    def test_provisioning_invalid_inputs_never_create_unintended_project(self):
        for email, password, role, project in [('bad', 'Private-password-long', 'po', 'p'),
            ('new@test.example', 'short', 'po', 'p'), ('new@test.example', 'Private-password-long', 'admin', 'p'),
            ('new@test.example', 'Private-password-long', 'po', 'missing')]:
            with self.subTest(email=email, role=role), self.assertRaises(AppError):
                provision_account(self.store, email, password, role, project)
        with self.store.db() as db: self.assertEqual(db.execute('SELECT COUNT(*) FROM users').fetchone()[0], 2)
        with self.assertRaises(AppError): existing_store(self.root / 'missing')
        self.assertFalse((self.root / 'missing').exists())

    def test_password_reset_during_login_cannot_mint_a_stale_session(self):
        import hashlib
        original = hashlib.pbkdf2_hmac
        raced = False
        def race(*args, **kwargs):
            nonlocal raced
            if not raced:
                raced = True
                change_account(self.store, 'po@test.example', password='Planner-new-private!')
            return original(*args, **kwargs)
        with patch('app.store.hashlib.pbkdf2_hmac', side_effect=race), self.assertRaises(AppError):
            self.store.login('po@test.example', 'Planner-initial-private!')
        with self.store.db() as db: self.assertEqual(db.execute('SELECT COUNT(*) FROM sessions').fetchone()[0], 0)

    def test_backup_restore_preserves_cipher_history_and_revokes_sessions(self):
        source = self.store.put('p', 'source', {'encrypted_text': self.store.encrypt(CANARY)})
        self.store.update('p', 'source', source['id'], {'encrypted_text': self.store.encrypt(CANARY + 'v2')})
        token, _ = self.store.login('owner@test.example', 'Owner-initial-private!')
        archive = self.root / 'archive.npdbackup'; phrase = 'Independent-backup-secret!'
        backup(self.store, archive, phrase)
        self.assertEqual(archive.stat().st_mode & 0o777, 0o600)
        self.assertNotIn(CANARY.encode(), archive.read_bytes())
        self.assertNotIn((self.store.directory / 'private.key').read_bytes(), archive.read_bytes())
        restored = self.root / 'restored'; restore(archive, restored, phrase)
        s = Store(restored)
        with self.assertRaises(AppError): s.authenticate(token)
        self.assertEqual(s.decrypt(s.get('p', 'source', source['id'])['encrypted_text']), CANARY + 'v2')
        self.assertEqual(s.decrypt(s.history('p', 'source', source['id'])[0]['encrypted_text']), CANARY)
        s.login('owner@test.example', 'Owner-initial-private!')
        original = self.store.path.read_bytes()
        with self.assertRaises(AppError): restore(archive, self.store.directory, phrase)
        self.assertEqual(self.store.path.read_bytes(), original)
        with self.assertRaises(AppError): backup(self.store, archive, phrase)

    def test_wrong_password_tampering_or_key_mismatch_never_install_partial_restore(self):
        self.store.put('p', 'source', {'encrypted_text': self.store.encrypt(CANARY)})
        archive = self.root / 'archive.npdbackup'; phrase = 'Independent-backup-secret!'
        backup(self.store, archive, phrase); original = archive.read_bytes()
        cases = [(original, 'Wrong-password-long!')]
        for seed in range(12):
            data = bytearray(original); at = random.Random(seed).randrange(len(data)); data[at] ^= 1
            cases.append((bytes(data), phrase))
        offset = len(MAGIC); cipher = backup_cipher(phrase, original[offset:offset+16])
        payload = json.loads(cipher.decrypt(original[offset+16:]))
        payload['files']['private.key'] = base64.b64encode(Fernet.generate_key()).decode()
        cases.append((original[:offset+16] + cipher.encrypt(json.dumps(payload).encode()), phrase))
        for index, (data, password) in enumerate(cases):
            with self.subTest(case=index):
                archive.write_bytes(data); dest = self.root / f'restore-{index}'
                with self.assertRaises(AppError): restore(archive, dest, password)
                self.assertFalse(dest.exists())
        self.assertEqual(list(self.root.glob('npd-restore-*')), [])


class RemoteGateTests(unittest.TestCase):
    def test_gate_conditions_remain_active_with_python_optimization(self):
        import subprocess
        import sys
        proc = subprocess.run([sys.executable, '-O', '-c', 'from scripts.live_validation import require; require(False)'], capture_output=True, text=True)
        self.assertNotEqual(proc.returncode, 0)
        self.assertIn('Workflow validation condition failed', proc.stderr)

    def test_https_url_contract_rejects_credential_paths_and_insecure_origins(self):
        self.assertEqual(https_origin('https://npd.example/'), 'https://npd.example')
        for url in ['http://npd.example', 'https://user:pass@npd.example', 'https://npd.example/projects',
                    'https://npd.example?key=secret', 'https://npd.example/#fragment', 'https://npd.example:bad']:
            with self.subTest(url=url), self.assertRaises(AppError): https_origin(url)

    def test_remote_harness_isolates_project_and_logs_out_with_explicit_test_provider(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp); model = RecordingModel(); model.provider = 'bedrock'
            model.probe = lambda: {'connection_verified': True}
            for role in ('owner', 'po'): store.create_user(role + '@example.test', role + '-private-test-password!', role, 'existing')
            server = Server(('127.0.0.1', 0), Service(store, model))
            thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
            try:
                result = run_remote(server.origin, {r: (r + '@example.test', r + '-private-test-password!') for r in ('owner', 'po')}, require_secure=False)
                self.assertEqual(result['status'], 'PASS', result)
                self.assertFalse(result['model_input_guarded'])
                self.assertFalse(result['secure_cookie_checked'])
                self.assertEqual(store.list('existing', 'source'), [])
                self.assertEqual(len(store.list(result['project_id'], 'source')), 1)
                with store.db() as db: self.assertEqual(db.execute('SELECT COUNT(*) FROM sessions').fetchone()[0], 0)
            finally:
                server.shutdown(); server.server_close(); thread.join(5)
