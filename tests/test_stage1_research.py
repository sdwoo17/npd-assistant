"""Synthetic regression coverage for structured Stage 1 and PRD handoff."""
import base64
import copy
import io
import json
import unittest
from unittest.mock import patch
from docx import Document
from app.contracts import validate_claims
from app.ingest import decode_file
from app.source_locations import document_locations
from app.research_contracts import TYPES, TASKS, PRD_SECTIONS
from app.store import AppError
from tests.helpers import Fixture, encoded, CANARY
from tests.test_stage2 import PlanningModel
from tests.research_factory import synthetic_fields, reference, item, pack, graph, scope_with_source


class ResearchModel(PlanningModel):
    def generate(self,task,payload,images=None):
        if task!='research_task':return super().generate(task,payload,images)
        self.calls.append((task,copy.deepcopy(payload)))
        if self.hook:self.hook()
        values=synthetic_fields(payload['output_type'])
        if payload['output_type']=='measurement':values.update(measurement_state='OBSERVED',value=999)
        return {'title':'Synthetic generated research','values':[{'key':k,'text':'\n'.join(v) if isinstance(v,list) else '' if v is None else str(v)} for k,v in values.items()]}


class Stage1ResearchTests(unittest.TestCase):
    def setUp(self):
        self.f=Fixture();self.s=self.f.service;self.u=self.f.po;self.s.model=ResearchModel()
        self.f.insight=self.s.knowledge(self.u['project_id'])[0]
    def tearDown(self):self.f.close()
    def post(self,path,body):return self.s.post(self.u,path,body)
    def brief(self,**fields):return item(self.s,self.u,'brief',fields=synthetic_fields('brief',**fields))
    def asset(self,purpose='existing_service'):
        return self.post('/api/planning-assets',encoded('synthetic.md',CANARY,title='Synthetic planning input',purpose=purpose))
    def review(self,row):return self.post('/api/research-workspace/items/review',{'item_id':row['id'],'expected_version':row['version'],'reason':'Synthetic review'})

    def test_review_replaces_secondary_sections_and_questions_atomically(self):
        asset=self.asset();row=self.post('/api/service-analysis',{'asset_refs':[reference(asset)],'prompt':'Synthetic baseline'})
        changed=self.post('/api/research-results/review',{'result_id':row['id'],'expected_version':row['version'],
            'text':'PO revised content','questions':['Revised open question'],'reason':'Removed stale interpretation'})
        self.assertEqual(changed['sections'][0]['text'],changed['text'])
        self.assertEqual(changed['questions'],['Revised open question'])
        self.assertNotIn(row['sections'][0]['text'],json.dumps(changed['sections']))
        with self.assertRaises(AppError):self.post('/api/research-results/review',{'result_id':row['id'],'expected_version':row['version'],'text':'stale'})

    def test_one_paragraph_prd_cannot_be_confirmed(self):
        row=self.post('/api/definitions',{'stage':'prd','title':'Incomplete','sections':[{'id':'overview','title':'Overview','text':'One line'}]})
        status=self.post('/api/definitions/readiness',{'document_id':row['id'],'expected_version':row['version']})
        self.assertFalse(status['ready']);self.assertEqual(len(status['coverage']),12)
        with self.assertRaises(AppError):self.post('/api/definitions/confirm',{'document_id':row['id'],'expected_version':row['version']})

    def test_all_task_contracts_reject_empty_review_and_na_requires_reason(self):
        brief=self.brief()
        self.assertEqual(len(TASKS),17)
        for task,(_,kind,_) in TASKS.items():
            if kind=='brief':continue
            row=item(self.s,self.u,kind,brief['id'],review=False,fields={k:None if s['type']=='number' else [] if s['type']=='lines' else False if s['type']=='checkbox' else '' for k,s in TYPES[kind].items()})
            with self.assertRaises(AppError,msg=task):self.review(row)
        row=item(self.s,self.u,'constraint',brief['id'],review=False,applicability='not_applicable')
        with self.assertRaises(AppError):self.review(row)
        row=self.post('/api/research-workspace/items',{'item_id':row['id'],'expected_version':row['version'],'na_reason':'Synthetic scope excludes deployment'})
        self.assertEqual(self.review(row)['state'],'reviewed')

    def test_explicit_scope_avoids_project_wide_context(self):
        first=self.brief(product_scope='Advertising');second=self.brief(product_scope='Shipping')
        a=item(self.s,self.u,'hypothesis',first['id'],title='Advertising only')
        b=item(self.s,self.u,'hypothesis',second['id'],title='Shipping only')
        with self.assertRaises(AppError):self.s.definition_context(self.u)
        selected=pack(self.s,self.u,first,[a])
        context=json.dumps(self.s.definition_context(self.u))
        self.assertIn(a['id'],context);self.assertNotIn(b['id'],context)
        with self.assertRaises(AppError):pack(self.s,self.u,first,[b])
        with self.assertRaises(AppError):self.s.post(self.f.other,'/api/research-workspace/select',{'pack_id':selected['id'],'expected_version':selected['version']})

    def test_numeric_meaning_and_unknown_baselines_are_guarded(self):
        evidence=[{'id':'INS-SYNTHETIC','text':'ROAS는 121%이다.','evidence_type':'research'}]
        with self.assertRaises(AppError):validate_claims('ROAS가 121% 개선되었다.',evidence)
        validate_claims('ROAS는 121%이다.',evidence)
        brief=self.brief()
        row=item(self.s,self.u,'numeric_claim',brief['id'],refs=[self.f.insight],review=False,
            fields=synthetic_fields('numeric_claim',source_kind='LEVEL',claim_kind='RELATIVE_CHANGE',value=121,unit='%'))
        with self.assertRaises(AppError):self.review(row)
        row=item(self.s,self.u,'metric',brief['id'],review=False,fields=synthetic_fields('metric',baseline_state='UNKNOWN',baseline=0))
        with self.assertRaises(AppError):self.review(row)

    def test_synthetic_or_research_insight_cannot_become_real_voc(self):
        brief=self.brief()
        for kind in ('REAL_VOC','INTERNAL_MEASUREMENT','SYNTHETIC_FGI'):
            row=item(self.s,self.u,'evidence',brief['id'],refs=[self.f.insight],review=False,
                fields=synthetic_fields('evidence',evidence_type=kind,supports_demand_validation=True))
            with self.assertRaises(AppError,msg=kind):self.review(row)

    def test_counts_preserve_real_synthetic_and_participant_denominators(self):
        brief=self.brief()
        row=item(self.s,self.u,'voc_collection',brief['id'],review=False,
            fields=synthetic_fields('voc_collection',utterances=3,participants=4,real_count=1,synthetic_count=2))
        with self.assertRaises(AppError):self.review(row)
        row=self.post('/api/research-workspace/items',{'item_id':row['id'],'expected_version':row['version'],
            'fields':synthetic_fields('voc_collection',utterances=3,participants=2,real_count=1,synthetic_count=2)})
        self.assertEqual(self.review(row)['fields']['real_count'],1)

    def test_actual_study_and_uat_need_execution_evidence(self):
        brief=self.brief()
        for kind,fields in [('actual_study',synthetic_fields('actual_study',run_state='EXECUTED')),
                            ('uat',synthetic_fields('uat',status='PASS'))]:
            row=item(self.s,self.u,kind,brief['id'],review=False,fields=fields)
            with self.assertRaises(AppError):self.review(row)
        planned=item(self.s,self.u,'actual_study',brief['id'])
        self.assertEqual(planned['fields']['run_state'],'NOT_RUN')

    def test_full_trace_supports_research_review_but_not_customer_certification(self):
        brief=self.brief();rows=graph(self.s,self.u,brief);selected=pack(self.s,self.u,brief,rows)
        result=self.post('/api/research-workspace/gate',{'pack_id':selected['id'],'expected_version':selected['version']})
        self.assertTrue(result['ready'],result['errors']);self.assertEqual(len(result['graph']),6)
        self.assertTrue(all(e['identity_check']=='CURRENT_VERSION' and e['semantic_review']=='PO_REVIEWED' and e['reviewed_by']==self.u['id'] for e in result['graph']))
        self.assertEqual(result['customer_validation'],'NOT_CERTIFIED')
        incomplete=pack(self.s,self.u,brief,rows[:-1]);result=self.post('/api/research-workspace/gate',{'pack_id':incomplete['id'],'expected_version':incomplete['version']})
        self.assertFalse(result['ready']);self.assertTrue(any('M' in e for e in result['errors']))

    def test_unresolved_conflict_is_kept_and_blocks_development(self):
        brief=self.brief();rows=graph(self.s,self.u,brief)
        other=item(self.s,self.u,'evidence',brief['id'])
        conflict=item(self.s,self.u,'conflict',brief['id'],refs=[rows[0],other])
        selected=pack(self.s,self.u,brief,rows+[other,conflict])
        current=self.s.pack_rows(self.u,selected)
        annotated=next(r for r in current if r['id']==rows[0]['id'])
        self.assertEqual(annotated['conflict_ids'],[conflict['id']])
        self.assertEqual(annotated['effective_claim_status'],'CONFLICTED')
        body={'pack_id':selected['id'],'expected_version':selected['version']}
        self.assertTrue(self.post('/api/research-workspace/gate',body)['ready'])
        self.assertFalse(self.post('/api/research-workspace/gate',{**body,'stage':'development'})['ready'])

    def test_prd_twelve_area_composition_and_confirmation_preserve_pack(self):
        brief=self.brief();rows=graph(self.s,self.u,brief);asset=self.asset()
        rows.append(item(self.s,self.u,'source_manifest',brief['id'],refs=[asset]))
        selected=pack(self.s,self.u,brief,rows)
        row=self.post('/api/research-workspace/compose',{'pack_id':selected['id'],'expected_version':selected['version'],'title':'Synthetic PRD'})
        self.assertEqual({s['id'] for s in row['sections']},set(PRD_SECTIONS))
        sections=[{**s,'text':s['text'] or 'Synthetic explicitly authored design section','coverage_status':'written'} for s in row['sections']]
        row=self.post('/api/definitions/update',{'document_id':row['id'],'expected_version':row['version'],'sections':sections,
            'questions':[{'text':'Actual baseline?','status':'unanswered','owner':'Synthetic analyst','next_action':'Collect approved data before development.'}]})
        confirmed=self.post('/api/definitions/confirm',{'document_id':row['id'],'expected_version':row['version']})
        self.assertEqual(confirmed['state'],'confirmed')
        exported=self.post('/api/definitions/export',{'document_id':confirmed['id'],'expected_version':confirmed['version']})
        self.assertEqual(exported['package']['research']['external_receipt'],'NOT_VERIFIED')
        self.assertEqual(exported['package']['research']['pack_ref']['id'],selected['id'])
        self.assertNotIn(asset['id'],json.dumps(exported));self.assertNotIn(CANARY,json.dumps(exported))

    def test_source_update_invalidates_pack_and_impact_redacts_bodies(self):
        brief=self.brief();asset=self.asset()
        row=item(self.s,self.u,'source_manifest',brief['id'],refs=[asset])
        selected=pack(self.s,self.u,brief,[row])
        self.post('/api/planning-assets/withdraw',{'asset_id':asset['id'],'expected_version':asset['version']})
        with self.assertRaises(AppError):self.s.selected_research(self.u)
        impact=self.s.get(self.u,'/api/research-workspace/impact')
        self.assertTrue(any(r['id']==selected['id'] for r in impact['affected']))
        self.assertNotIn(CANARY,json.dumps(self.s.get(self.u,'/api/research-workspace')))
        self.assertNotIn(row['title'],json.dumps(impact))

    def test_source_family_counts_do_not_treat_derivative_files_as_independent(self):
        brief=self.brief();rows=graph(self.s,self.u,brief)
        for language in ('Korean','English'):
            rows.append(item(self.s,self.u,'source_manifest',brief['id'],refs=[self.f.insight],
                fields=synthetic_fields('source_manifest',source_family_id='ONE_SYNTHETIC_FAMILY',language=language)))
        selected=pack(self.s,self.u,brief,rows)
        result=self.post('/api/research-workspace/gate',{'pack_id':selected['id'],'expected_version':selected['version']})
        self.assertEqual(result['source_family_count'],1)

    def test_export_delta_and_receipt_bind_hash_versions_and_current_permissions(self):
        brief=self.brief();row=item(self.s,self.u,'utility',brief['id'])
        first=pack(self.s,self.u,brief,[row]);a=self.post('/api/research-workspace/export',{'pack_id':first['id'],'expected_version':first['version']})
        changed=self.post('/api/research-workspace/items',{'item_id':row['id'],'expected_version':row['version'],'title':'Revised utility'})
        changed=self.review(changed);second=pack(self.s,self.u,brief,[changed])
        exported=self.post('/api/research-workspace/export',{'pack_id':second['id'],'expected_version':second['version'],'previous_export_id':a['export_id'],'reason':'Synthetic revision'})
        delta=next(c for c in exported['package']['changes'] if c['id']==row['id'])
        self.assertEqual(delta['before']['title'],row['title']);self.assertEqual(delta['after']['title'],'Revised utility')
        mapping=[{'source_id':r['id'],'source_version':r['version'],'destination_id':'destination-'+r['id']} for r in exported['package']['items']]
        body={'export_id':exported['export_id'],'package_hash':exported['hash'],'destination':'Synthetic receiver',
            'receipt_reference':'Synthetic receipt, not live AXIOM','mapping':mapping}
        with self.assertRaises(AppError):self.post('/api/research-workspace/receipt',{**body,'package_hash':'wrong'})
        with self.assertRaises(AppError):self.post('/api/research-workspace/receipt',{**body,'mapping':mapping[:-1]})
        receipt=self.post('/api/research-workspace/receipt',body)
        self.assertEqual(receipt['status'],'PO_RECONCILED');self.assertFalse(receipt['live_transport_verified'])
        self.post('/api/research-workspace/items/withdraw',{'item_id':changed['id'],'expected_version':changed['version']})
        with self.assertRaises(AppError):self.post('/api/research-workspace/receipt',body)

    def test_export_never_includes_original_asset_or_owner_content(self):
        brief=self.brief();asset=self.asset();row=item(self.s,self.u,'source_manifest',brief['id'],refs=[asset])
        selected=pack(self.s,self.u,brief,[row]);out=self.post('/api/research-workspace/export',{'pack_id':selected['id'],'expected_version':selected['version']})
        value=json.dumps(out)
        self.assertNotIn(CANARY,value);self.assertNotIn(asset['id'],value);self.assertNotIn('encrypted_',value)

    def test_ai_draft_is_scoped_and_cannot_certify_measurements(self):
        brief=self.brief();asset=self.asset()
        row=self.post('/api/research-workspace/generate',{'research_id':brief['id'],'output_type':'measurement','task_id':'VC-03',
            'prompt':'Synthetic measurement planning','input_refs':[reference(self.f.insight),reference(asset)]})
        self.assertEqual(row['state'],'draft');self.assertEqual(row['authorship'],'ai_proposed')
        self.assertEqual(row['fields']['measurement_state'],'UNKNOWN');self.assertIsNone(row['fields']['value'])
        self.assertNotIn(CANARY,json.dumps(self.s.model.calls));self.assertNotIn('encrypted_',json.dumps(self.s.model.calls))
        self.s.model.hook=lambda:self.s.store.update(self.u['project_id'],'insight',self.f.insight['id'],{'published':False},self.f.insight['version'])
        with self.assertRaises(AppError):self.post('/api/research-workspace/generate',{'research_id':brief['id'],'output_type':'measurement',
            'prompt':'Race','input_refs':[reference(self.s.knowledge(self.u['project_id'])[0])]})

    def test_docx_locations_preserve_paragraphs_and_table_cells(self):
        doc=Document();doc.add_paragraph('Synthetic paragraph');table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='Synthetic cell A';table.cell(0,1).text='Synthetic cell B'
        output=io.BytesIO();doc.save(output)
        body={'filename':'synthetic.docx','content_base64':base64.b64encode(output.getvalue()).decode()}
        _,content,_=decode_file(body);locations=document_locations(body,content)
        expected={'paragraph/1':'Synthetic paragraph','table/1/row/1/cell/1':'Synthetic cell A','table/1/row/1/cell/2':'Synthetic cell B'}
        self.assertEqual({r['locator']:content[r['start']:r['end']] for r in locations['parts']},expected)
        self.assertNotIn('Synthetic cell',json.dumps(locations))

    def test_untrusted_scope_fields_do_not_authorize_raw_research(self):
        brief=self.brief()
        with self.assertRaises(AppError):item(self.s,self.u,'source_manifest',brief['id'],refs=[self.f.source])
        with self.assertRaises(AppError):self.s.post(self.f.other,'/api/research-workspace/items',{'research_id':brief['id'],
            'output_type':'hypothesis','title':'Other project','fields':synthetic_fields('hypothesis'),'project_id':self.u['project_id']})

    def test_shared_ancestry_is_checked_once_per_access_traversal(self):
        brief=self.brief();rows=graph(self.s,self.u,brief);selected=pack(self.s,self.u,brief,rows)
        with patch.object(self.s,'knowledge',wraps=self.s.knowledge) as lookup:
            self.assertTrue(self.s.accessible(self.u['project_id'],selected))
        self.assertEqual(lookup.call_count,1)

    def test_proposed_input_limit_does_not_change_implemented_capabilities(self):
        brief=self.brief();f=synthetic_fields('input_policy',file_bytes=10*1024*1024,story_extractions=5,pixel_limit=16000000)
        row=item(self.s,self.u,'input_policy',brief['id'],fields=f)
        self.assertEqual(row['fields']['decision'],'PROPOSED')
        row=self.post('/api/research-workspace/items',{'item_id':row['id'],'expected_version':row['version'],'fields':{**f,'decision':'CONFIRMED_CURRENT'}})
        with self.assertRaises(AppError):self.review(row)
        self.assertEqual(self.s.get(self.u,'/api/research-workspace')['input_capabilities']['story_extractions'],3)


    def test_insight_partial_edit_preserves_synthetic_type_and_source_locations(self):
        source=self.s.store.get(self.u['project_id'],'source',self.f.source['id'])
        row=self.s.post(self.f.owner,'/api/insights',{'source_id':source['id'],'title':'Synthetic assumption',
            'text':'Synthetic assumption only','evidence_type':'SYNTHETIC_FGI','source_family_id':'Synthetic family',
            'source_locator':{'parts':source['source_locations']['parts']}})
        changed=self.s.post(self.f.owner,'/api/insights/update',{'insight_id':row['id'],'expected_version':row['version'],'title':'Edited synthetic assumption'})
        self.assertEqual(changed['evidence_type'],'synthetic')
        self.assertEqual(changed['evidence_category'],'SYNTHETIC_FGI')
        self.assertEqual(changed['source_locator'],row['source_locator'])
        self.assertEqual(changed['source_family_id'],'Synthetic family')



    def test_benchmark_cannot_supply_our_observed_baseline(self):
        brief=self.brief()
        f=synthetic_fields('metric',baseline_state='OBSERVED',baseline=121)
        row=item(self.s,self.u,'metric',brief['id'],refs=[self.f.insight],fields=f,review=False)
        with self.assertRaises(AppError):self.review(row)
        external=item(self.s,self.u,'measurement',brief['id'],refs=[self.f.insight],
            fields=synthetic_fields('measurement',measurement_state='OBSERVED',measurement_origin='EXTERNAL',value=121))
        row=item(self.s,self.u,'metric',brief['id'],refs=[external],fields=f,review=False)
        with self.assertRaises(AppError):self.review(row)
        measured=item(self.s,self.u,'measurement',brief['id'],refs=[self.asset()],
            fields=synthetic_fields('measurement',measurement_state='OBSERVED',value=121))
        row=item(self.s,self.u,'metric',brief['id'],refs=[measured],fields=f)
        self.assertEqual(row['state'],'reviewed')

    def test_alternative_comparison_and_unmeasured_utility_require_honest_status(self):
        brief=self.brief()
        row=item(self.s,self.u,'alternative',brief['id'],review=False,
            fields=synthetic_fields('alternative',compared_options=['Synthetic only option']))
        with self.assertRaises(AppError):self.review(row)
        row=self.post('/api/research-workspace/items',{'item_id':row['id'],'expected_version':row['version'],
            'fields':{**row['fields'],'fewer_options_reason':'Only one feasible option in this synthetic scenario.'}})
        self.assertEqual(self.review(row)['state'],'reviewed')
        row=item(self.s,self.u,'utility',brief['id'],review=False,fields=synthetic_fields('utility',before_minutes=0))
        with self.assertRaises(AppError):self.review(row)

    def test_legacy_pack_locator_is_preserved_without_certifying_offsets(self):
        row=self.s.post(self.f.owner,'/api/insights',{'source_id':self.f.source['id'],
            'title':'Synthetic legacy pack','text':'Synthetic legacy finding','source_locator':'page-002'})
        self.assertEqual(row['source_locator'],{'parts':[],'legacy_label':'page-002','status':'NEEDS_REVIEW'})
        edited=self.s.post(self.f.owner,'/api/insights/update',{'insight_id':row['id'],
            'expected_version':row['version'],'title':'Edited legacy finding'})
        self.assertEqual(edited['source_locator'],row['source_locator'])

if __name__=='__main__':unittest.main()
